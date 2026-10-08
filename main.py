import os, io, tempfile, requests
from fastapi import FastAPI, Request
from pptx import Presentation
from pptx.util import Inches

app = FastAPI()

# Acepta tus nombres actuales
WHATSAPP_TOKEN = os.getenv("WHATSAPP_TOKEN")
WHATSAPP_PHONE_ID = os.getenv("WHATSAPP_PHONE_ID") or os.getenv("PHONE_NUMBER_ID") or os.getenv("PHONE_ID")
VERIFY_TOKEN = os.getenv("VERIFY_TOKEN") or os.getenv("VERIFY_TOKEN")
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY") or os.getenv("OPENAI_KEY") or os.getenv("OPENAI1_KEY")

SESSIONS = {}

def send_text(to, text):
    url = f"https://graph.facebook.com/v20.0/{WHATSAPP_PHONE_ID}/messages"
    headers = {"Authorization": f"Bearer {WHATSAPP_TOKEN}", "Content-Type": "application/json"}
    data = {"messaging_product": "whatsapp", "to": to, "type": "text", "text": {"body": text}}
    requests.post(url, headers=headers, json=data)

def send_doc(to, pptx_bytes, filename="Reporte.pptx"):
    url_media = f"https://graph.facebook.com/v20.0/{WHATSAPP_PHONE_ID}/media"
    headers = {"Authorization": f"Bearer {WHATSAPP_TOKEN}"}
    files = {
        'file': (filename, pptx_bytes, 'application/vnd.openxmlformats-officedocument.presentationml.presentation'),
        'type': (None, 'application/vnd.openxmlformats-officedocument.presentationml.presentation'),
        'messaging_product': (None, 'whatsapp')
    }
    r = requests.post(url_media, headers=headers, files=files)
    if r.status_code!= 200:
        print("Error media:", r.text)
        send_text(to, f"Error subiendo PPTX: {r.text[:400]}")
        return
    media_id = r.json()["id"]
    url_msg = f"https://graph.facebook.com/v20.0/{WHATSAPP_PHONE_ID}/messages"
    headers_json = {"Authorization": f"Bearer {WHATSAPP_TOKEN}", "Content-Type": "application/json"}
    data = {
        "messaging_product": "whatsapp",
        "to": to,
        "type": "document",
        "document": {"id": media_id, "filename": filename, "caption": "Listo ✅ Reporte generado"}
    }
    requests.post(url_msg, headers=headers_json, json=data)

def create_pptx(texts, images):
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(6)).text_frame.text = "\n".join(texts) if texts else "Reporte de mantenimiento"
    for img_b in images:
        try:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as tf:
                tf.write(img_b); path = tf.name
            slide.shapes.add_picture(path, Inches(0.5), Inches(0.5), Inches(12), Inches(6.5))
            os.unlink(path)
        except Exception as e:
            print("img error", e)
    bio = io.BytesIO(); prs.save(bio); bio.seek(0); return bio.getvalue()

@app.get("/")
def home():
    return {"status": "OK", "phone_id_loaded": bool(WHATSAPP_PHONE_ID), "phone_id": WHATSAPP_PHONE_ID}

@app.get("/webhook")
def verify(hub_verify_token: str = None, hub_challenge: str = None):
    if hub_verify_token == (os.getenv("VERIFY_TOKEN") or "automatyco123"):
        return int(hub_challenge)
    return "Token invalido"

@app.post("/webhook")
async def webhook(req: Request):
    data = await req.json()
    try:
        value = data["entry"][0]["changes"][0]["value"]
        if "messages" not in value: return "ok"
        msg = value["messages"][0]
        from_num = msg["from"]
        SESSIONS.setdefault(from_num, {"texts": [], "images": []})

        if msg["type"] == "text":
            body = msg["text"]["body"].strip()
            if body.lower() in ["generar reporte", "generar", "reporte"]:
                sess = SESSIONS[from_num]
                if not sess["texts"] and not sess["images"]:
                    send_text(from_num, "Aún no tengo datos. Mándame fotos primero."); return "ok"
                send_text(from_num, f"Generando reporte con {len(sess['images'])} fotos...")
                pptx = create_pptx(sess["texts"], sess["images"])
                send_doc(from_num, pptx, filename=f"Reporte_{from_num}.pptx")
                SESSIONS[from_num] = {"texts": [], "images": []}
            else:
                SESSIONS[from_num]["texts"].append(body)
                send_text(from_num, f"Guardado ({len(SESSIONS[from_num]['texts'])}). Escribe *generar reporte*")
        elif msg["type"] == "image":
            mid = msg["image"]["id"]
            h = {"Authorization": f"Bearer {WHATSAPP_TOKEN}"}
            info = requests.get(f"https://graph.facebook.com/v20.0/{mid}", headers=h).json()
            img_data = requests.get(info["url"], headers=h).content
            SESSIONS[from_num]["images"].append(img_data)
            send_text(from_num, f"Foto {len(SESSIONS[from_num]['images'])} guardada. Escribe *generar reporte*")
    except Exception as e:
        print("Error:", e)
    return "ok"
